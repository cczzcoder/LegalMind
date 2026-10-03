"""第 1 层解析器单元测试（设计 §6、§7、§14.2）：不需要数据库。

覆盖接口契约、各解析器、分块与 §6 定位不变量、注册表分发、内存守卫与资源上限。
PDF 用 ``tests.helpers.build_minimal_pdf`` 手工构造的最小 PDF 作样本（无外部依赖、内容确定），
真实法律文本的基准见 ``scripts/benchmark_parsers.py``。
"""

import io

import docx
import pytest

from app.modules.documents.validation import DOCX, HTML, PDF, TEXT
from app.modules.parsing.chunking import chunk_blocks
from app.modules.parsing.interface import (
    QUALITY_NEEDS_REVIEW,
    QUALITY_OK,
    DocumentParser,
    DocumentTooLarge,
    MemoryBudgetExceeded,
    NormalizedTextBuilder,
    ParseError,
    ParseLimits,
    UnsupportedFormat,
    clean_text,
    normalize_bbox,
)
from app.modules.parsing.memory import MemoryGuard, current_rss_bytes, peak_working_set_bytes
from app.modules.parsing.native.docx import DocxParser
from app.modules.parsing.native.html import HtmlParser
from app.modules.parsing.native.pdf import PdfplumberPdfParser, Pypdfium2PdfParser
from app.modules.parsing.native.text import TextParser
from app.modules.parsing.registry import LAYER3_PARSERS, PDF_BACKENDS, get_parser, parse_document
from tests.helpers import build_minimal_pdf as build_pdf

GENEROUS = ParseLimits(
    max_bytes=10 * 1024 * 1024,
    max_pages=100,
    max_chars=1_000_000,
    memory_budget_bytes=512 * 1024 * 1024,
)


# --------------------------------------------------------------------------- 样本构造


def build_docx() -> bytes:
    document = docx.Document()
    document.add_heading("第一章　总则", level=1)
    document.add_paragraph("第一条　正文一段。")
    document.add_paragraph("")  # 空段落应被跳过
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "甲栏"
    table.cell(0, 1).text = "乙栏"
    document.add_paragraph("表后段落。")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def assert_location_invariants(result) -> None:
    """所有解析器都必须满足的 §6 定位不变量。"""
    # 规范化文本就是各块用 \n 连接（chunk 还原与 text_sha256 都依赖这一点）
    assert result.text == "\n".join(block.text for block in result.blocks)
    for block in result.blocks:
        assert block.char_start >= 0
        assert block.char_end > block.char_start
        assert result.text[block.char_start : block.char_end] == block.text
        # 与 chunk_spans 的 CHECK 约束一致：坐标系与 bbox 必须同时有或同时无
        assert (block.coordinate_system is None) == (block.bbox is None)
        if block.bbox is not None:
            x0, y0, x1, y1 = block.bbox
            assert 0.0 <= x0 <= x1 <= 1.0
            assert 0.0 <= y0 <= y1 <= 1.0
    # 块在规范化文本中首尾相接，块间恰好一个 \n
    for left, right in zip(result.blocks, result.blocks[1:]):
        assert left.char_end + 1 == right.char_start


# --------------------------------------------------------------------------- 接口基础


def test_normalized_text_builder_tracks_offsets():
    builder = NormalizedTextBuilder()
    assert builder.add("   ") is None  # 空块被跳过
    first = builder.add("第一条　甲")
    second = builder.add("第二条　乙\n")
    assert (first.start, first.end) == (0, 5)
    assert (second.start, second.end) == (6, 11)
    assert builder.build() == "第一条　甲\n第二条　乙"
    assert builder.length == 11


def test_clean_text_normalizes_line_endings_and_whitespace():
    assert clean_text("  a \r\n b \r ") == "a\n b"
    assert clean_text("\r\n") == ""


@pytest.mark.parametrize(
    ("box", "expected"),
    [
        ((0.0, 0.0, 100.0, 200.0), (0.0, 0.0, 0.5, 0.5)),
        ((200.0, 100.0, 0.0, 0.0), (0.0, 0.0, 1.0, 0.25)),  # 顺序颠倒也能纠正
        ((-10.0, -10.0, 210.0, 410.0), (0.0, 0.0, 1.0, 1.0)),  # 越界裁剪
    ],
)
def test_normalize_bbox(box, expected):
    assert normalize_bbox(*box, 200.0, 400.0) == pytest.approx(expected)


def test_normalize_bbox_rejects_degenerate_page():
    assert normalize_bbox(0, 0, 1, 1, 0, 100) is None
    assert normalize_bbox(0, 0, 1, 1, 100, -1) is None


# --------------------------------------------------------------------------- 分块


def test_chunk_round_trip_reproduces_text():
    result = parse_document(build_docx(), DOCX)
    chunks = chunk_blocks(result.blocks, result.text)
    assert [chunk.ordinal for chunk in chunks] == list(range(len(chunks)))
    assert "\n".join(chunk.text for chunk in chunks) == result.text
    for chunk in chunks:
        assert chunk.text
        for index, span in enumerate(chunk.spans):
            assert span.ordinal == index
            assert span.block.char_end > span.block.char_start


def test_chunk_keeps_blocks_whole_and_respects_target_size():
    result = parse_document(build_docx(), DOCX)
    chunks = chunk_blocks(result.blocks, result.text, max_chars=20)
    assert len(chunks) > 1
    # 除超长单块外，每个 chunk 不超过目标长度；块永不被切开
    covered = [span.block for chunk in chunks for span in chunk.spans]
    assert covered == list(result.blocks)
    for chunk in chunks:
        if len(chunk.spans) > 1:
            # chunk 文本 = 各块文本 + 块间 \n
            assert len(chunk.text) <= 20 + len(chunk.spans) - 1


def test_chunk_handles_empty_document():
    assert chunk_blocks((), "") == []


# --------------------------------------------------------------------------- 纯文本


def test_text_parser_lines_and_offsets():
    result = TextParser().parse("第一行\r\n\r\n第二行\n".encode(), media_type=TEXT, limits=GENEROUS)
    assert result.media_type == TEXT
    assert [block.text for block in result.blocks] == ["第一行", "第二行"]
    assert result.pages == ()
    assert_location_invariants(result)


def test_text_parser_enforces_limits():
    with pytest.raises(DocumentTooLarge):
        TextParser().parse(b"x" * 100, media_type=TEXT, limits=_limits(max_bytes=10))
    with pytest.raises(DocumentTooLarge):
        TextParser().parse(b"x" * 100, media_type=TEXT, limits=_limits(max_chars=10))


def _limits(**overrides) -> ParseLimits:
    base = {
        "max_bytes": 10 * 1024 * 1024,
        "max_pages": 100,
        "max_chars": 1_000_000,
        "memory_budget_bytes": 512 * 1024 * 1024,
    }
    base.update(overrides)
    return ParseLimits(**base)


# --------------------------------------------------------------------------- HTML


def test_html_parser_extracts_leaf_blocks_and_drops_script():
    markup = (
        "<html><body>"
        "<h1>标题</h1>"
        "<div><p>第一段<em>强调</em>。</p><p>第二段</p></div>"
        "<script>var x = 1;</script><style>p{color:red}</style>"
        "<ul><li>甲</li><li>乙</li></ul>"
        "</body></html>"
    )
    result = HtmlParser().parse(markup.encode(), media_type=HTML, limits=GENEROUS)
    assert [block.text for block in result.blocks] == ["标题", "第一段强调。", "第二段", "甲", "乙"]
    assert result.blocks[0].kind == "heading" and result.blocks[0].level == 1
    assert result.blocks[3].kind == "list_item"
    # 无页概念的格式不写页码与 bbox（设计 §6）
    assert result.pages == ()
    assert all(block.page_index is None and block.bbox is None for block in result.blocks)
    assert_location_invariants(result)


def test_html_parser_decodes_entities_and_utf8():
    markup = "<html><body><p>甲 &amp; 乙 &lt;条款&gt;</p></body></html>"
    result = HtmlParser().parse(markup.encode(), media_type=HTML, limits=GENEROUS)
    assert result.blocks[0].text == "甲 & 乙 <条款>"


def test_html_parser_rejects_non_utf8():
    with pytest.raises(ParseError):
        HtmlParser().parse("<p>甲</p>".encode("gbk"), media_type=HTML, limits=GENEROUS)


def test_html_parser_rejects_unparsable_input():
    with pytest.raises(ParseError):
        HtmlParser().parse(b"\x00\x01\x02", media_type=HTML, limits=GENEROUS)


# --------------------------------------------------------------------------- DOCX


def test_docx_parser_keeps_body_order_and_skips_empty():
    result = DocxParser().parse(build_docx(), media_type=DOCX, limits=GENEROUS)
    assert [block.text for block in result.blocks] == [
        "第一章　总则",
        "第一条　正文一段。",
        "甲栏",
        "乙栏",
        "表后段落。",
    ]
    assert result.blocks[0].kind == "heading" and result.blocks[0].level == 1
    assert result.blocks[2].kind == "table_cell"
    assert result.pages == ()
    assert_location_invariants(result)


def test_docx_parser_rejects_invalid_package():
    with pytest.raises(ParseError):
        DocxParser().parse(b"not a docx at all", media_type=DOCX, limits=GENEROUS)


# --------------------------------------------------------------------------- PDF 双后端


@pytest.mark.parametrize("parser_class", [Pypdfium2PdfParser, PdfplumberPdfParser])
def test_pdf_parsers_produce_location(parser_class):
    content = build_pdf([["Hello Legal Mind", "Second line here"], ["Page two text"]])
    result = parser_class().parse(content, media_type=PDF, limits=GENEROUS)
    assert result.page_count == 2
    assert [block.text for block in result.blocks] == [
        "Hello Legal Mind",
        "Second line here",
        "Page two text",
    ]
    assert [block.page_index for block in result.blocks] == [0, 0, 1]
    assert all(block.coordinate_system == "normalized-page" for block in result.blocks)
    assert_location_invariants(result)


def test_pdf_backends_agree_on_text_and_structure():
    content = build_pdf([["Hello Legal Mind", "Second line here"], ["Page two text"]])
    left = Pypdfium2PdfParser().parse(content, media_type=PDF, limits=GENEROUS)
    right = PdfplumberPdfParser().parse(content, media_type=PDF, limits=GENEROUS)
    assert left.text == right.text
    assert [block.text for block in left.blocks] == [block.text for block in right.blocks]
    assert left.page_count == right.page_count


def test_pdf_without_text_layer_is_marked_needs_review():
    result = Pypdfium2PdfParser().parse(build_pdf([[]]), media_type=PDF, limits=GENEROUS)
    assert result.char_count == 0
    assert result.quality_status == QUALITY_NEEDS_REVIEW
    assert result.warnings == ("no_text_layer",)
    assert result.pages[0].has_text is False


def test_pdf_text_layer_is_ok_quality():
    result = Pypdfium2PdfParser().parse(build_pdf([["Hello"]]), media_type=PDF, limits=GENEROUS)
    assert result.quality_status == QUALITY_OK
    assert result.warnings == ()


def test_pdf_enforces_max_pages():
    content = build_pdf([["one"], ["two"], ["three"]])
    with pytest.raises(DocumentTooLarge):
        Pypdfium2PdfParser().parse(content, media_type=PDF, limits=_limits(max_pages=2))


def test_pdf_enforces_max_bytes():
    with pytest.raises(DocumentTooLarge):
        Pypdfium2PdfParser().parse(
            build_pdf([["Hello"]]), media_type=PDF, limits=_limits(max_bytes=16)
        )


@pytest.mark.parametrize("parser_class", [Pypdfium2PdfParser, PdfplumberPdfParser])
def test_pdf_rejects_invalid_bytes(parser_class):
    with pytest.raises(ParseError):
        parser_class().parse(b"%PDF-1.7\ngarbage", media_type=PDF, limits=GENEROUS)


def test_pdf_memory_budget_aborts_parse(monkeypatch):
    """PDF 解析必须逐页调用内存守卫；守卫本身的判定用真实分配另测。"""
    from app.modules.parsing import memory

    calls = {"count": 0}

    def fake_rss() -> int:
        calls["count"] += 1
        return 0 if calls["count"] == 1 else 10**9  # 第一次取基线，之后必然超预算

    monkeypatch.setattr(memory, "current_rss_bytes", fake_rss)

    content = build_pdf([["Hello Legal Mind"], ["Second page"]])
    with pytest.raises(MemoryBudgetExceeded):
        Pypdfium2PdfParser().parse(
            content, media_type=PDF, limits=_limits(memory_budget_bytes=1024)
        )
    assert calls["count"] >= 2


# --------------------------------------------------------------------------- 注册表


def test_registry_dispatches_by_media_type():
    assert isinstance(get_parser(PDF), Pypdfium2PdfParser)
    assert isinstance(get_parser(PDF, pdf_backend="pdfplumber"), PdfplumberPdfParser)
    assert isinstance(get_parser(DOCX), DocxParser)
    assert isinstance(get_parser(HTML), HtmlParser)
    assert isinstance(get_parser(TEXT), TextParser)


def test_registry_rejects_unknown_media_type_and_backend():
    with pytest.raises(UnsupportedFormat):
        get_parser("application/msword")
    with pytest.raises(UnsupportedFormat):
        get_parser(PDF, pdf_backend="docling")


def test_docling_layer3_is_deferred_not_registered():
    assert "docling" in LAYER3_PARSERS
    assert "docling" not in PDF_BACKENDS
    assert not any(
        isinstance(parser_class(), DocumentParser) and "docling" in parser_class.name
        for parser_class in PDF_BACKENDS.values()
    )


# --------------------------------------------------------------------------- 内存守卫


def test_memory_guard_raises_when_budget_exceeded():
    with pytest.raises(MemoryBudgetExceeded), MemoryGuard(1024 * 1024) as guard:
        blob = bytearray(64 * 1024 * 1024)
        blob[::4096] = b"\x01" * len(blob[::4096])
        guard.check()
    assert current_rss_bytes() > 0


def test_memory_guard_disabled_when_budget_not_positive():
    with MemoryGuard(0) as guard:
        guard.check()  # 不限制时不抛
    assert guard.peak_growth_bytes >= 0


def test_peak_working_set_is_monotonic():
    first = peak_working_set_bytes()
    assert first >= 0
    assert peak_working_set_bytes() >= first
